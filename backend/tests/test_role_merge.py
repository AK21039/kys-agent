import json

import pytest

from kys import config
from kys.ontology import EntityType, MergeDecision, NodeKind, ReviewStatus
from kys.process import role_merge
from kys.process.roles import RoleIndex
from kys.textnorm import fold

ROL = EntityType.ROL.value


def _ent(conn, name, etype=ROL):
    return conn.execute("INSERT INTO entities(entity_type, canonical_name, normalized_name) VALUES (?,?,?)",
                        (etype, name, fold(name))).lastrowid


@pytest.fixture
def kb(conn):
    doc = conn.execute(
        """INSERT INTO documents(code, code_key, title, doc_type, body_text, file_path, file_sha256, meta_source)
           VALUES ('PR-01', 'pr1', 'Bakım', 'PROSEDUR', '-', 'x.docx', 'sha', 'regex')""").lastrowid
    ids = {n: _ent(conn, n) for n in ("Mekanik Bakım Teknisyeni", "Mekanik Bakım Teknisyenleri",
                                      "Genel Müdür", "Genel Müdür Yardımcısı", "Satınalma Müdürü",
                                      "Satın Alma Müdürü", "Kalite Müdürü")}
    for subj in ("Mekanik Bakım Teknisyeni", "Mekanik Bakım Teknisyeni", "Mekanik Bakım Teknisyenleri",
                 "Satınalma Müdürü", "Genel Müdür"):
        conn.execute("INSERT INTO rules(document_id, rule_type, statement, subject_entity_id, quote) "
                     "VALUES (?, 'SORUMLULUK', 's', ?, 'q')", (doc, ids[subj]))
    conn.execute("INSERT INTO rules(document_id, rule_type, statement, object_entity_id, quote) "
                 "VALUES (?, 'ONAY', 's', ?, 'q')", (doc, ids["Satın Alma Müdürü"]))
    pid = conn.execute("INSERT INTO processes(name, owner_entity_id) VALUES ('P', ?)",
                       (ids["Satın Alma Müdürü"],)).lastrowid
    conn.execute("INSERT INTO process_steps(process_id, seq, name, responsible_entity_id, approver_entity_id) "
                 "VALUES (?, 1, 'a', ?, ?)", (pid, ids["Mekanik Bakım Teknisyenleri"], ids["Satın Alma Müdürü"]))
    conn.execute("INSERT INTO decisions(title, owner_entity_id, approval_chain) VALUES ('K', ?, ?)",
                 (ids["Satın Alma Müdürü"], json.dumps([ids["Satın Alma Müdürü"], ids["Genel Müdür"]])))
    # iki varyant da aynı kenarı taşıyor → birleşince tekil kalmalı
    for name in ("Satınalma Müdürü", "Satın Alma Müdürü"):
        conn.execute("INSERT INTO edges(src_kind, src_id, rel_type, dst_kind, dst_id) VALUES (?, ?, 'X', ?, 1)",
                     (NodeKind.VARLIK.value, ids[name], NodeKind.DOKUMAN.value))
    conn.execute("INSERT INTO entity_aliases(entity_id, alias, normalized_alias) VALUES (?, 'Satınalma Md.', ?)",
                 (ids["Satın Alma Müdürü"], fold("Satınalma Md.")))
    conn.commit()
    return conn, ids


def _names(groups):
    return [sorted(m.name for m in g.members) for g in groups]


def test_suggest_groups_no_chaining(kb):
    conn, _ = kb
    groups = role_merge.suggest_groups(conn)
    names = _names(groups)
    assert sorted(["Mekanik Bakım Teknisyeni", "Mekanik Bakım Teknisyenleri"]) in names
    assert sorted(["Satınalma Müdürü", "Satın Alma Müdürü"]) in names
    assert not any("Genel Müdür Yardımcısı" in g and "Genel Müdür" in g for g in names)
    assert not any("Kalite Müdürü" in g for g in names)
    teknisyen = next(g for g in groups if "Mekanik" in g.canonical.name)
    assert teknisyen.canonical.name == "Mekanik Bakım Teknisyeni"   # en çok kullanılan ana ad önerilir


def test_xlsx_roundtrip_and_rejection_memory(kb):
    conn, _ = kb
    groups = role_merge.suggest_groups(conn)
    data = role_merge.render_xlsx(groups)
    from io import BytesIO

    from openpyxl import load_workbook
    wb = load_workbook(BytesIO(data))
    ws = wb["Rol Grupları"]
    for row in ws.iter_rows(min_row=2):
        if row[3].value.startswith("Mekanik"):
            row[6].value = MergeDecision.EVET.value
        elif row[3].value == "Satınalma Müdürü":
            row[6].value = MergeDecision.EVET.value
            row[7].value = "Satınalma Müdürü"
        elif row[3].value == "Satın Alma Müdürü":
            row[6].value = MergeDecision.HAYIR.value
    buf = BytesIO()
    wb.save(buf)
    decisions = role_merge.read_decisions(buf.getvalue())
    by_ana = {d.ana: d for d in decisions}
    assert set(by_ana["Mekanik Bakım Teknisyeni"].evet) == {"Mekanik Bakım Teknisyeni", "Mekanik Bakım Teknisyenleri"}
    assert by_ana["Satınalma Müdürü"].hayir == ["Satın Alma Müdürü"]

    role_merge.save_decisions(decisions)
    saved = json.loads(config.settings.role_merges_path.read_text(encoding="utf-8"))
    assert saved[ROL]["ayri"] == [sorted([fold("Satın Alma Müdürü"), fold("Satınalma Müdürü")])]
    assert len(saved[ROL]["birlesik"]) == 1   # tek üyeli satınalma grubu birleştirme sayılmaz
    # reddedilen çift bir daha aynı gruba önerilmez
    assert sorted(["Satınalma Müdürü", "Satın Alma Müdürü"]) not in _names(role_merge.suggest_groups(conn))


def test_merge_moves_every_reference(kb):
    conn, ids = kb
    keep, gone = ids["Satınalma Müdürü"], ids["Satın Alma Müdürü"]
    s = role_merge.merge_group(conn, ROL, "Satınalma Müdürü", ["Satın Alma Müdürü"])
    conn.commit()
    assert s["birlesen"] == 1
    assert conn.execute("SELECT COUNT(*) FROM entities WHERE id = ?", (gone,)).fetchone()[0] == 0
    for table, col in role_merge._FK_COLUMNS:
        assert conn.execute(f"SELECT COUNT(*) FROM {table} WHERE {col} = ?", (gone,)).fetchone()[0] == 0
    assert conn.execute("SELECT object_entity_id FROM rules WHERE rule_type = 'ONAY'").fetchone()[0] == keep
    chain = json.loads(conn.execute("SELECT approval_chain FROM decisions").fetchone()[0])
    assert chain == [keep, ids["Genel Müdür"]]
    assert conn.execute("SELECT COUNT(*) FROM edges WHERE src_id = ?", (keep,)).fetchone()[0] == 1
    aliases = {r[0] for r in conn.execute("SELECT alias FROM entity_aliases WHERE entity_id = ?", (keep,))}
    assert aliases == {"Satın Alma Müdürü", "Satınalma Md."}
    assert conn.execute("SELECT review_status FROM entities WHERE id = ?", (keep,)).fetchone()[0] == \
        ReviewStatus.ONAYLANDI.value
    m = RoleIndex(conn).resolve("Satınalma Md.")
    assert (m.entity_id, m.method) == (keep, "takma_ad")


def test_rename_to_new_canonical_and_idempotent_reapply(kb):
    conn, ids = kb
    data = {ROL: {"birlesik": [{"ana": "Mekanik Teknisyen",
                                "varyantlar": ["Mekanik Bakım Teknisyeni", "Mekanik Bakım Teknisyenleri"]}],
                  "ayri": []}}
    first = role_merge.apply_saved(conn, data)
    assert first["birlesen"] == 1
    row = conn.execute("SELECT id, canonical_name FROM entities WHERE normalized_name = ?",
                       (fold("Mekanik Teknisyen"),)).fetchone()
    assert row["id"] == ids["Mekanik Bakım Teknisyeni"]   # en çok kullanılan varlık korunur, adı değişir
    assert conn.execute("SELECT COUNT(*) FROM rules WHERE subject_entity_id = ?", (row["id"],)).fetchone()[0] == 3
    again = role_merge.apply_saved(conn, data)
    assert again == {"grup": 0, "birlesen": 0, "tasinan": 0, "takma_ad": 0}


def test_preview_does_not_write(kb):
    conn, ids = kb
    data = {ROL: {"birlesik": [{"ana": "Satınalma Müdürü", "varyantlar": ["Satın Alma Müdürü"]}], "ayri": []}}
    s = role_merge.apply_saved(conn, data, commit=False)
    assert s["birlesen"] == 1
    assert conn.execute("SELECT COUNT(*) FROM entities WHERE id = ?", (ids["Satın Alma Müdürü"],)).fetchone()[0] == 1


def test_rebuilt_db_gets_saved_merges_after_reimport(kb):
    """DB yeniden kurulup varyant yeniden açılsa bile JSON'daki karar tekrar uygulanır."""
    conn, ids = kb
    role_merge.save_decisions([role_merge.Decision(ROL, "Satınalma Müdürü", ["Satın Alma Müdürü"])])
    role_merge.apply_saved(conn)
    again = _ent(conn, "Satın Alma Müdürü")   # yeni çıkarım eski yazımı tekrar açtı
    conn.commit()
    s = role_merge.apply_saved(conn)
    assert s["birlesen"] == 1
    assert conn.execute("SELECT COUNT(*) FROM entities WHERE id = ?", (again,)).fetchone()[0] == 0


def test_bad_decision_value_is_error(kb):
    conn, _ = kb
    from io import BytesIO

    from openpyxl import load_workbook
    wb = load_workbook(BytesIO(role_merge.render_xlsx(role_merge.suggest_groups(conn))))
    wb["Rol Grupları"]["G2"] = "belki"
    buf = BytesIO()
    wb.save(buf)
    with pytest.raises(ValueError, match="KARAR"):
        role_merge.read_decisions(buf.getvalue())
