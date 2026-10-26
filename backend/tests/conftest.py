from pathlib import Path

import pytest

from kys import db


@pytest.fixture(autouse=True)
def isolated_overrides(tmp_path, monkeypatch):
    """Testler gerçek data/kys_overrides.json dosyasına asla dokunmaz."""
    import dataclasses

    from kys import config
    monkeypatch.setattr(config, "settings",
                        dataclasses.replace(config.settings, overrides_path=tmp_path / "overrides.json"))


@pytest.fixture
def conn(tmp_path):
    c = db.connect(tmp_path / "test.db")
    db.init_db(c)
    yield c
    c.close()


def _numbered(paragraph):
    """Paragrafa Word otomatik numaralandırması ekler (metne numara yazılmaz)."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    ppr = paragraph._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    for tag, val in (("w:ilvl", "0"), ("w:numId", "1")):
        el = OxmlElement(tag)
        el.set(qn("w:val"), val)
        num_pr.append(el)
    ppr.append(num_pr)


def make_docx(path: Path, code: str = "PR-04", rev: str = "03", title: str = "Satınalma Prosedürü",
              body: list[tuple[str, str]] | None = None, header_layout: str = "rows") -> Path:
    """Türk KYS dokümanı benzeri DOCX üretir: üstbilgi tablosu, başlıklar, gövde tablosu."""
    import docx
    from docx.shared import Cm

    d = docx.Document()
    header = d.sections[0].header
    if header_layout == "rows":
        t = header.add_table(rows=2, cols=4, width=Cm(16))
        for i, v in enumerate(["Doküman No", "Yayın Tarihi", "Rev. No", "Rev. Tarihi"]):
            t.cell(0, i).text = v
        for i, v in enumerate([code, "01.01.2020", rev, "05.06.2024"]):
            t.cell(1, i).text = v
    else:
        t = header.add_table(rows=1, cols=4, width=Cm(16))
        for i, v in enumerate(["Doküman No", code, "Revizyon No", rev]):
            t.cell(0, i).text = v
    header.add_paragraph(title.replace("i", "İ").replace("ı", "I").upper())  # Türkçe büyük harf

    body = body or [
        ("h1", "AMAÇ"),
        ("p", "Bu prosedürün amacı satınalma faaliyetlerinin kurallarını belirlemektir."),
        ("h1", "UYGULAMA"),
        ("h2", "4.1 Talep"),
        ("p", "Satınalma talebi FR-12 Satınalma Talep Formu ile yapılır."),
        ("p", "1. Talep formu doldurulur."),
        ("h2", "4.2 Onay"),
        ("p", "Satınalma talebi 100.000 TL'yi aşarsa Genel Müdür onaylar. Tedarikçi seçimi için PR-08'e bakınız."),
        ("table", "Tutar|Onaycı;0-100.000 TL|Satınalma Müdürü;100.000 TL üzeri|Genel Müdür"),
    ]
    for kind, text in body:
        if kind == "h1":
            _numbered(d.add_heading(text, level=1))
        elif kind == "h2":
            d.add_heading(text, level=2)
        elif kind == "table":
            rows = [r.split("|") for r in text.split(";")]
            t = d.add_table(rows=len(rows), cols=len(rows[0]))
            for i, row in enumerate(rows):
                for j, v in enumerate(row):
                    t.cell(i, j).text = v
        else:
            d.add_paragraph(text)
    d.save(str(path))
    return path


def make_xlsx(path: Path) -> Path:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Yetki Matrisi"
    ws.append(["Doküman No", "YM-01", "Rev. No", "02"])
    ws.append([])
    ws.append(["Karar", "Satınalma Müdürü", "Genel Müdür"])
    ws.append(["Satınalma onayı ≤100.000 TL", "Onaylar", "Bilgilendirilir"])
    ws.append(["Tedarikçi sözleşmesi feshi", None, "Onaylar"])
    wb.save(str(path))
    return path
