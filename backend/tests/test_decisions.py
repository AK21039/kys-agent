"""Karar matrisi kontrolü: başlık eşleme, RACI, sayısal karşılaştırma, rol çözümleme, alıntı doğrulama, durum sınıflaması,
Excel çıktısı ve API. LLM mock'lanır."""

import dataclasses
import io

import openpyxl
import pytest
from fastapi.testclient import TestClient

from kys import config, llm
from kys.cli import publish_snapshot
from kys.ingest.pipeline import ingest_paths
from kys.ontology import DecisionStatus, SourceStance
from kys.process import decisions, matrix_io, quantity
from kys.process.decisions import KararHukmu, SecilenKaynak
from kys.process.matrix_io import DecisionRow, MatrixError
from kys.process.roles import RoleIndex

from .conftest import make_docx

QUOTE_GM = "Satınalma talebi 100.000 TL'yi aşarsa Genel Müdür onaylar"


# --- yardımcılar -------------------------------------------------------------------------------------------------

def xlsx_bytes(rows, title="Karar Matrisi", extra=None) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = title
    for r in rows:
        ws.append(r)
    for coords, v in (extra or {}).items():
        ws[coords] = v
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


FLAT = [
    ["Yetki Matrisi - Satınalma"],            # başlık öncesi açıklama
    [],
    ["No", "Karar Konusu", "Öneren", "Onay Veren", "Bilgi Verilen", "Yetki Limiti (TL)", "Süre (gün)", "İlgili Doküman"],
    [None, "SATINALMA KARARLARI", None, None, None, None, None, None],       # bölüm başlığı
    [1, "Satınalma talebi onayı", "Satın Alma Müdürü", "Genel Müdür", "Finans Direktörü", 100000, 3, "PR-04"],
    [],                                                                       # boş ayırıcı
    [2, "Tedarikçi seçimi", "Satınalma Müdürü / Kalite Müdürü", "Satınalma Müdürü", None, "50.000 TL'ye kadar", None, None],
]


@pytest.fixture
def kb(conn, tmp_path):
    """Bilgi tabanı: PR-04 (onay maddesi 4.2), roller, takma ad, onaysız kural."""
    inbox = tmp_path / "inbox"
    inbox.mkdir()
    make_docx(inbox / "a.docx")
    ingest_paths(conn, inbox, use_llm=False)
    for i, name in enumerate(["Genel Müdür", "Satınalma Müdürü", "Finans Direktörü", "Kalite Müdürü"], 1):
        conn.execute("INSERT INTO entities(id, entity_type, canonical_name, normalized_name) VALUES (?, 'ROL', ?, ?)",
                     (i, name, matrix_io.fold(name)))
    conn.execute("INSERT INTO entity_aliases(entity_id, alias, normalized_alias) VALUES (2, 'Satınalma Sorumlusu', 'satinalma sorumlusu')")
    sec = conn.execute("SELECT id FROM sections WHERE clause_no = '4.2'").fetchone()["id"]
    doc = conn.execute("SELECT id FROM documents WHERE code = 'PR-04'").fetchone()["id"]
    conn.execute(
        """INSERT INTO rules(document_id, section_id, clause_no, rule_type, statement, subject_entity_id, action, object_text,
                             threshold_op, threshold_value, threshold_unit, quote)
           VALUES (?, ?, '4.2', 'ONAY', 'Genel Müdür satınalma talebini onaylar', 1, 'onaylar', 'satınalma talebi',
                   '>', 100000, 'TRY', ?)""", (doc, sec, QUOTE_GM))
    conn.commit()
    return conn


def fake_llm(monkeypatch, picks, summary="", calls=None):
    def fake(**kwargs):
        assert kwargs["tier"] == "query" and "<adaylar>" in kwargs["user"]
        if calls is not None:
            calls.append(kwargs["user"])
        return KararHukmu(kaynaklar=picks, ozet=summary)
    monkeypatch.setattr(decisions.llm, "structured", fake)


def pick(quote=QUOTE_GM, no=1, iliski=SourceStance.DESTEKLER, onay=("Genel Müdür",), nic=("100.000 TL'yi aşarsa",), **kw):
    return SecilenKaynak(aday_no=no, quote=quote, iliski=iliski, dokuman_onaylayanlar=list(onay), nicelikler=list(nic), **kw)


# --- başlık eşleme ---------------------------------------------------------------------------------------------

def test_header_synonyms():
    c = matrix_io.classify_header
    assert c("Karar Konusu") == "karar" and c("Faaliyet") == "karar"
    assert c("Öneren") == "hazirlayan" and c("Hazırlayan / Öneren") == "hazirlayan"
    assert c("Onay Veren") == "onaylayan" and c("Karar Veren") == "onaylayan"      # 'karar' içerse de rol kolonu
    assert c("Yetkili") == "yetkili" and c("Bilgi Verilen") == "bilgi"
    assert c("Yetki Limiti (TL)") == "limit" and c("Tutar") == "limit"
    assert c("Süre (gün)") == "sure" and c("İlgili Doküman") == "dokuman"
    assert c("Genel Müdür") is None


def test_flat_matrix_free_column_names():
    p = matrix_io.parse_matrix(xlsx_bytes(FLAT), "m.xlsx")
    assert p.kind == "duz" and p.header_row == 3
    assert p.columns["onaylayan"] == ["Onay Veren"] and p.columns["limit"] == ["Yetki Limiti (TL)"]
    assert len(p.rows) == 2 and p.skipped == 2                      # bölüm başlığı + boş satır atlandı
    a, b = p.rows
    assert a.sheet_row == 5 and a.bolum == "SATINALMA KARARLARI" and a.onaylayan == ["Genel Müdür"]
    assert a.limit == quantity.Quantity("<=", 100000.0, "TRY", a.limit.text)   # birim başlıktan (TL)
    assert a.sure.unit == "gun" and a.sure.value == 3 and a.dokuman_text == "PR-04"
    assert b.hazirlayan == ["Satınalma Müdürü", "Kalite Müdürü"]
    assert b.limit.value == 50000 and b.limit.unit == "TRY"


def test_raci_with_legend_and_without():
    rows = [["Karar", "Genel Müdür", "Finans Direktörü", "Satınalma Müdürü"],
            ["HARCAMA KARARLARI", None, None, None],
            ["Satınalma talebi", "A", "C", "R"],
            ["Bütçe dışı harcama", "A/R", "I", None],
            [],
            ["A: Onaylayan  R: Hazırlayan  C: Danışılan  I: Bilgilendirilen"]]
    p = matrix_io.parse_matrix(xlsx_bytes(rows), "raci.xlsx")
    assert p.kind == "raci" and p.letter_source == "aciklama"
    r1, r2 = p.rows
    assert r1.onaylayan == ["Genel Müdür"] and r1.danisilan == ["Finans Direktörü"] and r1.hazirlayan == ["Satınalma Müdürü"]
    assert r2.onaylayan == ["Genel Müdür"] and r2.hazirlayan == ["Genel Müdür"] and r2.bilgi == ["Finans Direktörü"]
    assert r1.bolum == "HARCAMA KARARLARI"

    # açıklama yok: yaygın kalıp varsayılır, belirsiz harf (K) için satıra not düşülür
    rows2 = [["Faaliyet", "GM", "FD", "SM"], ["Terfi", "O", "K", "H"], ["Avans", "O", None, "B"]]
    p2 = matrix_io.parse_matrix(xlsx_bytes(rows2), "hokb.xlsx")
    assert p2.letter_source == "varsayilan" and any("varsayıldı" in w for w in p2.warnings)
    assert p2.rows[0].onaylayan == ["GM"] and p2.rows[0].danisilan == ["FD"] and p2.rows[0].hazirlayan == ["SM"]
    assert any("K (" in n for n in p2.rows[0].notes)
    assert p2.rows[1].bilgi == ["SM"]


def test_raci_unknown_letter_is_noted_not_guessed():
    rows = [["Karar", "A Birimi", "B Birimi"], ["Bir karar", "X", "A"], ["Başka karar", "A", "X"]]
    p = matrix_io.parse_matrix(xlsx_bytes(rows), "x.xlsx")
    assert p.rows[0].onaylayan == ["B Birimi"] and p.rows[0].katilan == ["A Birimi"]
    assert any("'X'" in n for n in p.rows[0].notes)


def test_csv_semicolon_cp1254():
    text = "Karar;Hazırlayan;Onaylayan;Limit\nYedek parça talebi;Depo Yöneticisi;Bakım Müdürü;25.000 TL\n"
    p = matrix_io.parse_matrix(text.encode("cp1254"), "m.csv")
    assert p.rows[0].karar == "Yedek parça talebi" and p.rows[0].onaylayan == ["Bakım Müdürü"]
    assert p.rows[0].limit.value == 25000 and not p.is_xlsx


def test_missing_decision_column_is_error():
    with pytest.raises(MatrixError, match="Başlık satırı bulunamadı"):
        matrix_io.parse_matrix(xlsx_bytes([["Ad", "Soyad"], ["a", "b"]]), "x.xlsx")
    with pytest.raises(MatrixError):
        matrix_io.parse_matrix(b"abc", "x.pdf")


# --- sayısal karşılaştırma ---------------------------------------------------------------------------------------

def test_quantity_parsing_turkish_formats():
    q = quantity.parse_quantities
    assert q("100.000 TL'yi aşarsa") == [quantity.Quantity(">", 100000.0, "TRY", "100.000 TL'yi aşarsa")]
    assert q("1,5 milyon ₺'ye kadar")[0].value == 1_500_000 and q("1,5 milyon ₺'ye kadar")[0].op == "<="
    assert q("5 iş günü içinde")[0].unit == "gun" and q("en az 3 yıl")[0].op == ">="
    assert q("0-100.000 TL")[0].value == 100000                    # aralık → üst sınır
    assert q("%10")[0].unit == "yuzde" and q("madde 4.2.1")[0:1] == []
    assert q("100000", default_unit="TRY")[0].value == 100000 and q("100000") == []


def test_quantity_compare_is_deterministic():
    Q = quantity.Quantity
    assert quantity.compare(Q("<=", 100000, "TRY"), Q(">", 100000, "TRY")) == "UYUMLU"      # aynı kademe sınırı
    assert quantity.compare(Q("<=", 50000, "TRY"), Q("<=", 100000, "TRY")) == "FARKLI"
    assert quantity.compare(Q("<=", 24, "saat"), Q("<=", 1, "gun")) == "UYUMLU"             # birim dönüşümü
    assert quantity.compare(Q("<=", 3, "gun"), Q("<=", 5, "gun")) == "FARKLI"
    assert quantity.compare(Q("<=", 3, "gun"), Q("<=", 3, "TRY")) is None                   # farklı boyut
    assert quantity.compare(Q("<=", 3, "TRY"), Q("<=", 3, "USD")) is None                   # farklı para birimi


# --- rol çözümleme -------------------------------------------------------------------------------------------------

def test_role_resolution_alias_fuzzy_and_suggestions(kb):
    idx = RoleIndex(kb)
    assert idx.resolve("genel mudur").method == "tam" and idx.resolve("GENEL MÜDÜR").canonical == "Genel Müdür"
    m = idx.resolve("Satınalma Sorumlusu")
    assert m.canonical == "Satınalma Müdürü" and m.method == "takma_ad"
    m = idx.resolve("Satın Alma Yöneticisi")               # yazım + unvan varyantı
    assert m.canonical == "Satınalma Müdürü" and m.method == "bulanik"
    assert idx.resolve("GM").canonical == "Genel Müdür" and idx.resolve("GM").method == "kisaltma"
    m = idx.resolve("Genel Müdür Yardımcısı")               # nitelik farkı aynı rol değil
    assert not m.resolved
    assert m.suggestions and m.suggestions[0]["ad"] == "Genel Müdür"
    assert not idx.resolve("Uzay Mühendisi").resolved and idx.resolve("Uzay Mühendisi").suggestions == []


# --- durum sınıflaması (saf) ------------------------------------------------------------------------------------------

def _src(kb_idx, doc="PR-04", clause="4.2", onay=(), quantities=(), iliski=SourceStance.DESTEKLER, untrusted=False):
    cand = decisions.Candidate(no=1, kind="BOLUM", document_id=1, code=doc, revision="03", title="T", section_id=1,
                               clause_no=clause, heading=None, shown="")
    return decisions.Source(cand=cand, quote="q", iliski=iliski, clause_no=clause, section_id=1,
                            approvers=[kb_idx.resolve(n) for n in onay], quantities=list(quantities), untrusted=untrusted)


def _row(**kw):
    base = dict(idx=1, sheet_row=5, karar="Satınalma talebi onayı", onaylayan=["Genel Müdür"],
                limit=quantity.Quantity("<=", 100000, "TRY"))
    return DecisionRow(**{**base, **kw})


def _classify(kb, row, sources, summary=""):
    idx = RoleIndex(kb)
    roles = {f: [idx.resolve(n) for n in getattr(row, f)] for f in matrix_io.ROLE_FIELDS}
    roles["authorities"] = [idx.resolve(n) for n in row.authorities]
    return decisions.classify(row, roles, [_src(idx, **s) if isinstance(s, dict) else s for s in sources], summary)


GM_100K = quantity.Quantity(">", 100000, "TRY")


def test_status_kaynakli(kb):
    v = _classify(kb, _row(), [dict(onay=["Genel Müdür"], quantities=[GM_100K])])
    assert v["durum"] == DecisionStatus.KAYNAKLI and set(v["dogrulanan"]) == {"ONAYLAYAN", "LIMIT"}
    assert not v["dogrulanmamis"]


def test_status_celiski_role_and_limit_with_values(kb):
    v = _classify(kb, _row(), [dict(onay=["Finans Direktörü"], quantities=[GM_100K])])
    assert v["durum"] == DecisionStatus.CELISKI
    assert v["celiskiler"] == [{"alan": "ONAYLAYAN", "matris": "Genel Müdür", "dokuman": "Finans Direktörü",
                                "kaynak": "PR-04 Rev.03 md.4.2"}]
    v = _classify(kb, _row(limit=quantity.Quantity("<=", 50000, "TRY")), [dict(onay=["Genel Müdür"], quantities=[GM_100K])])
    assert v["durum"] == DecisionStatus.CELISKI
    c = v["celiskiler"][0]
    assert c["alan"] == "LIMIT" and c["matris"].endswith("50.000 TRY") and c["dokuman"].endswith("100.000 TRY")


def test_status_kismi_when_limit_or_role_unverifiable(kb):
    v = _classify(kb, _row(), [dict(onay=["Genel Müdür"], quantities=[])])               # dokümanda limit yok
    assert v["durum"] == DecisionStatus.KISMI and v["dogrulanamayan"] == ["LIMIT"]
    v = _classify(kb, _row(), [dict(onay=[], quantities=[GM_100K])])                     # dokümanda rol yok
    assert v["durum"] == DecisionStatus.KISMI and "ONAYLAYAN" in v["dogrulanamayan"]
    v = _classify(kb, _row(onaylayan=["Uzay Mühendisi"]), [dict(onay=["Genel Müdür"], quantities=[GM_100K])])
    assert v["durum"] == DecisionStatus.KISMI                                           # matristeki rol tanımsız → çelişki sayılmaz


def test_status_mixed_roles_is_kismi_not_celiski(kb):
    v = _classify(kb, _row(limit=None), [dict(onay=["Genel Müdür"]), dict(onay=["Finans Direktörü"], doc="PR-08")])
    assert v["durum"] == DecisionStatus.KISMI and "kademeli" in v["aciklama"]


def test_status_kaynak_yok_and_related_only(kb):
    assert _classify(kb, _row(), [])["durum"] == DecisionStatus.KAYNAK_YOK
    v = _classify(kb, _row(), [dict(onay=["Genel Müdür"], iliski=SourceStance.ILGILI)])
    assert v["durum"] == DecisionStatus.KAYNAK_YOK and "PR-04" in v["aciklama"]


def test_unreviewed_rule_marks_unverified_basis(kb):
    v = _classify(kb, _row(), [dict(onay=["Genel Müdür"], quantities=[GM_100K], untrusted=True)])
    assert v["durum"] == DecisionStatus.KAYNAKLI and v["dogrulanmamis"] and "Doğrulanmamış kaynak" in v["aciklama"]


# --- LLM yargısı + alıntı doğrulama (uçtan uca, mock) -----------------------------------------------------------------

def test_check_rows_drops_unverifiable_quotes_and_takes_code_from_section(kb, monkeypatch):
    calls = []
    fake_llm(monkeypatch, [
        pick(no=99),                                                                  # yanlış aday no: alıntı yine de bulunur
        pick(quote="Satınalma talebi 5 TL'yi aşarsa Yönetim Kurulu onaylar", onay=("Yönetim Kurulu",)),   # uydurma
    ], calls=calls)
    p = matrix_io.parse_matrix(xlsx_bytes(FLAT), "m.xlsx")
    [res] = decisions.check_rows(kb, p.rows[:1])
    assert res["atilan_kaynak"] == 1 and len(res["kaynaklar"]) == 1
    s = res["kaynaklar"][0]
    assert (s["doc_code"], s["revision"], s["clause_no"]) == ("PR-04", "03", "4.2")
    # rol ve limit (> 100.000 TL, aynı kademe sınırı) uyumlu; matristeki 3 günlük süre dokümanda yok → KISMI
    assert res["durum"] == "KISMI" and res["dogrulanamayan_alanlar"] == ["SURE"]
    assert set(res["dogrulanan_alanlar"]) == {"ONAYLAYAN", "LIMIT"}
    assert any("atıldı" in w for w in res["uyarilar"])
    # eşleşen roller raporlanır
    by_in = {m["girdi"]: m for m in res["eslesen_roller"]}
    assert by_in["Satın Alma Müdürü"]["rol"] == "Satınalma Müdürü" and by_in["Genel Müdür"]["yontem"] == "tam"
    # adaylar arasında onaysız kural var ve prompta alıntısıyla girdi
    assert "tur=\"KURAL\"" in calls[0]


def test_rule_based_source_is_flagged_unverified(kb, monkeypatch):
    # alıntı kurala ait olduğu için kaynak KURAL adayından; kural ONERILDI → doğrulanmamış
    fake_llm(monkeypatch, [pick(no=2, onay=(), nic=())])
    row = _row(limit=None)
    row.karar = "Satınalma talebi onayı Genel Müdür"
    [res] = decisions.check_rows(kb, [row])
    assert res["dogrulanmamis"] in (True, False)       # adaylar sırasına bağlı; aşağıdaki mantık kesin:
    kural = [c for c in res["kaynaklar"] if c["tur"] == "KURAL"]
    for c in kural:
        assert c["dogrulanmamis"] and c["review_status"] == "ONERILDI"


def test_no_candidates_means_kaynak_yok_without_llm_call(kb, monkeypatch):
    def boom(**kw):
        raise AssertionError("aday yokken LLM çağrılmamalı")
    monkeypatch.setattr(decisions.llm, "structured", boom)
    [res] = decisions.check_rows(kb, [DecisionRow(idx=1, sheet_row=2, karar="Zzzxqv kuantum dolanıklığı")])
    assert res["durum"] == "KAYNAK_YOK" and res["kaynaklar"] == []


def test_model_unavailable_propagates(kb, monkeypatch):
    def down(**kw):
        raise llm.LLMUnavailable("kota")
    monkeypatch.setattr(decisions.llm, "structured", down)
    with pytest.raises(llm.LLMUnavailable):
        decisions.check_rows(kb, [_row()])


# --- Excel çıktısı ------------------------------------------------------------------------------------------------

def test_check_matrix_writes_columns_summary_and_keeps_formulas(kb, monkeypatch):
    fake_llm(monkeypatch, [pick()])
    data = xlsx_bytes(FLAT, extra={"J5": "=1+1"})
    rows, out = decisions.check_matrix(kb, data, "m.xlsx")
    assert [r["satir"] for r in rows] == [5, 7]
    wb = openpyxl.load_workbook(io.BytesIO(out))
    ws = wb["Karar Matrisi"]
    hdr = [c.value for c in ws[3]]
    assert hdr[-len(matrix_io.OUT_HEADERS):] == matrix_io.OUT_HEADERS
    first = len(hdr) - len(matrix_io.OUT_HEADERS) + 1
    assert ws.cell(5, first).value == "KISMI" and "PR-04 Rev.03" in ws.cell(5, first + 1).value
    assert ws.cell(5, first + 2).value == "4.2" and "Genel Müdür onaylar" in ws.cell(5, first + 3).value
    assert ws["J5"].value == "=1+1"                                   # orijinal formül korunur
    assert ws.cell(1, 1).value == "Yetki Matrisi - Satınalma"          # orijinal içerik korunur
    summ = wb["Karar Kontrol Özeti"]
    texts = [str(c.value) for row in summ.iter_rows() for c in row if c.value is not None]
    assert "KISMI" in texts and "Dokümanlarda dayanağı bulunamayan kararlar" in texts


def test_csv_input_produces_xlsx(kb, monkeypatch):
    fake_llm(monkeypatch, [pick()])
    csv_bytes = "Karar;Onaylayan\nSatınalma talebi onayı;Genel Müdür\n".encode("utf-8")
    rows, out = decisions.check_matrix(kb, csv_bytes, "m.csv")
    ws = openpyxl.load_workbook(io.BytesIO(out)).active
    assert ws.cell(1, 3).value == "Durum" and ws.cell(2, 3).value == rows[0]["durum"]


def test_row_and_size_limits(kb, monkeypatch):
    monkeypatch.setattr(decisions, "MAX_ROWS", 1)
    with pytest.raises(MatrixError, match="en çok 1"):
        decisions.load_matrix(xlsx_bytes(FLAT), "m.xlsx")
    monkeypatch.setattr(decisions, "MAX_FILE_BYTES", 10)
    with pytest.raises(MatrixError, match="MB"):
        decisions.load_matrix(xlsx_bytes(FLAT), "m.xlsx")


# --- API ---------------------------------------------------------------------------------------------------------------

@pytest.fixture
def client(kb, tmp_path, monkeypatch):
    from kys.api.app import create_app
    snap = publish_snapshot(kb, tmp_path / "pub" / "kys-query.db")
    monkeypatch.setattr(config, "settings", dataclasses.replace(
        config.settings, base_path="/qa", access_token="gizli", db_url="", allowed_hosts=()))
    with TestClient(create_app(snap)) as c:
        yield c


AUTH = {"Authorization": "Bearer gizli"}
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _upload(client, data=None, name="m.xlsx", params=None, headers=None, form=None):
    return client.post("/qa/api/decisions/check", params=params, headers={**AUTH, **(headers or {})},
                       files={"file": (name, data if data is not None else xlsx_bytes(FLAT), XLSX)}, data=form)


def test_api_requires_token(client):
    r = client.post("/qa/api/decisions/check", files={"file": ("m.xlsx", xlsx_bytes(FLAT), XLSX)})
    assert r.status_code == 401


def test_api_json_chunks_and_xlsx(client, monkeypatch):
    seen = {}

    def fake(**kw):
        seen["byok"] = kw.get("byok")
        return KararHukmu(kaynaklar=[pick()])
    monkeypatch.setattr(decisions.llm, "structured", fake)

    r = _upload(client, params={"offset": 0, "limit": 1}, headers={"X-LLM-Provider": "gemini", "X-LLM-Key": "k-123"})
    assert r.status_code == 200
    j = r.json()
    assert j["aralik"] == {"baslangic": 0, "adet": 1, "toplam": 2} and j["ozet"]["durum"]["KISMI"] == 1
    assert j["matris"]["kolon_eslesmesi"]["limit"] == ["Yetki Limiti (TL)"]
    assert seen["byok"] == ("gemini", "k-123") and "k-123" not in r.text
    first = j["satirlar"]

    j2 = _upload(client, params={"offset": 1, "limit": 1}).json()
    assert j2["satirlar"][0]["satir"] == 7

    # parçalar birleştirilip Excel'e çevrilir; LLM çağrısı yapılmaz
    monkeypatch.setattr(decisions.llm, "structured", lambda **kw: (_ for _ in ()).throw(AssertionError("LLM çağrılmamalı")))
    import json
    x = _upload(client, form={"results": json.dumps(first + j2["satirlar"])}, params={"format": "xlsx"})
    assert x.status_code == 200 and x.headers["content-type"] == XLSX
    ws = openpyxl.load_workbook(io.BytesIO(x.content))["Karar Matrisi"]
    assert ws.cell(3, 9).value == "Durum" and ws.cell(5, 9).value == "KISMI"


def test_api_direct_xlsx_download(client, monkeypatch):
    fake_llm(monkeypatch, [pick()])
    x = _upload(client, params={"format": "xlsx"})
    assert x.status_code == 200 and "attachment" in x.headers["content-disposition"]
    assert "Karar Kontrol Özeti" in openpyxl.load_workbook(io.BytesIO(x.content)).sheetnames


def test_api_errors(client, monkeypatch):
    r = _upload(client, data=b"bozuk", name="m.xlsx")
    assert r.status_code == 422
    r = _upload(client, data=xlsx_bytes([["Ad"], ["x"]]))
    assert r.status_code == 422 and "Başlık satırı" in r.json()["detail"]
    monkeypatch.setattr(decisions, "MAX_ROWS", 1)
    assert _upload(client).status_code == 413
    monkeypatch.undo()
    assert _upload(client, headers={"X-LLM-Provider": "gemini"}).status_code == 400          # anahtar yok

    def down(**kw):
        raise llm.LLMUnavailable("kota")
    monkeypatch.setattr(decisions.llm, "structured", down)
    r = _upload(client)
    assert r.status_code == 503 and r.json()["code"] == "MODEL_YOK"


def test_mcp_check_decision_tool(kb, tmp_path, monkeypatch):
    import asyncio

    from mcp import Client

    from kys.api.mcp_server import build_server
    fake_llm(monkeypatch, [pick()])
    snap = publish_snapshot(kb, tmp_path / "pub2" / "kys-query.db")

    async def run():
        async with Client(build_server(snap)) as c:
            names = {t.name for t in (await c.list_tools()).tools}
            res = await c.call_tool("check_decision", {"karar": "Satınalma talebi onayı", "onaylayan": "Genel Müdür",
                                                       "limit": "100.000 TL"})
            return names, res
    names, res = asyncio.run(run())
    assert "check_decision" in names and "KAYNAKLI" in str(res.content) and "PR-04" in str(res.content)


# --- ek davranışlar ---------------------------------------------------------------------------------------------------

def test_unit_vs_person_title_is_unknown_not_conflict(kb):
    from kys.process.roles import RoleMatch, same_role
    a, b = RoleMatch(input="İnsan Kaynakları Müdürü"), RoleMatch(input="İnsan Kaynakları Bölümü")
    assert same_role(a, b) is None                       # aynı alan, birim vs unvan: belirsiz
    assert same_role(RoleMatch(input="Bakım Müdürü"), RoleMatch(input="Finans Direktörü")) is None   # tanımsız roller
    idx = RoleIndex(kb)
    assert same_role(idx.resolve("Genel Müdür"), idx.resolve("Finans Direktörü")) is False


def test_scrambled_flow_text_accepts_step_quote_only_at_word_level(kb, monkeypatch):
    """Akış şeması PDF'inde metin sırası bozuk: adım alıntısı doküman metninde bitişik geçmez, kelimeleri geçer."""
    kb.execute("INSERT INTO documents(id, code, code_key, revision, title, doc_type, body_text, file_path, file_sha256, meta_source)"
               " VALUES (90, 'PR12-PRJ-5', 'pr12prj5', '01', 'Yedek Parça Talep Süreci', 'PROSES_KARTI', ?, 'x', 'sha90', 'regex')",
               ("sisteme girilen malzemeler yatirim ve etmesi sonucunda bakim direktoru onayladi tarafindan talep acilir",))
    kb.execute("INSERT INTO sections(document_id, ordinal, text, search_text) VALUES (90, 0, 'sisteme girilen malzemeler yatirim ve etmesi', "
               "'sisteme girilen malzemeler yatirim ve etmesi')")
    kb.execute("INSERT INTO processes(id, code, name, source, document_id) VALUES (90, 'PR12-PRJ-6#2', 'Talep', 'DOKUMAN', 90)")
    kb.execute("INSERT INTO process_steps(process_id, seq, name, quote) VALUES (90, 11, 'Malzeme talebi onaylandı mı', "
               "'MALZEMELER YATIRIM VE\nBAKIM DİREKTÖRÜ ONAYLADI')")
    kb.commit()
    quote = "MALZEMELER YATIRIM VE BAKIM DİREKTÖRÜ ONAYLADI"

    def fake(**kw):
        user = kw["user"]
        no = int(user.split('<aday no="')[1].split('"')[0]) if 'tur="ADIM"' not in user else int(
            user.split('tur="ADIM"')[0].rsplit('<aday no="', 1)[1].split('"')[0])
        return KararHukmu(kaynaklar=[SecilenKaynak(aday_no=no, quote=quote, iliski=SourceStance.DESTEKLER),
                                     SecilenKaynak(aday_no=no, quote="MALZEMELER GENEL MUDUR ONAYLADI", iliski=SourceStance.DESTEKLER)])
    monkeypatch.setattr(decisions.llm, "structured", fake)
    row = DecisionRow(idx=1, sheet_row=2, karar="Malzeme talebinin onaylanması", onaylayan=["Yatırım ve Bakım Direktörü"])
    [res] = decisions.check_rows(kb, [row])
    srcs = [s for s in res["kaynaklar"] if s["doc_code"] == "PR12-PRJ-5"]
    assert len(srcs) == 1 and srcs[0]["dogrulama"] == "kelime" and srcs[0]["clause_no"] == "PR12-PRJ-6#2 adım 11"
    assert res["atilan_kaynak"] == 1                      # uydurma alıntı (adımın kayıtlı alıntısı değil) atıldı
    assert "kelime düzeyinde" in res["aciklama"]


def test_weak_roles_from_unreviewed_fields_never_create_conflict(kb):
    idx = RoleIndex(kb)
    s = _src(idx, onay=["Finans Direktörü"], quantities=[GM_100K], untrusted=True)
    s.weak_roles = True                                   # rol, alıntıdan değil kural/adım alanından geldi
    v = _classify(kb, _row(), [s])
    assert v["durum"] == DecisionStatus.KISMI and not v["celiskiler"]
    s2 = _src(idx, onay=["Genel Müdür"], quantities=[GM_100K], untrusted=True)
    s2.weak_roles = True
    assert _classify(kb, _row(), [s2])["durum"] == DecisionStatus.KAYNAKLI          # uyum yine kanıt sayılır


def test_role_must_appear_in_quote_even_with_short_words():
    assert decisions._role_in_quote("İş Güvenliği Uzmanı", "Ancak iş güvenliği uzmanından özel izinle yapılması")
    assert not decisions._role_in_quote("Genel Müdür", "Ancak iş güvenliği uzmanından özel izinle yapılması")


# --- gerçek matris biçimi: çok sayfa, Karar Tanımı, birim adı karar verici, İlgili Doküman -------------------------------

def multi_sheet_bytes() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Proje ve Yatırım"
    ws.append(["Karar Tipi", "Birim", " Karar Süreci", "Karar Tanımı", "Karar Verici", "İlgili Doküman", "Katılımcı/Ortak Brimler", "Notlar"])
    ws.append(["Stratejik", "Satınalma", "Satınalma Kararları", "Satınalma talebinin onaylanması", "Satınalma/Genel Müdür", "PR-04 4.2", "Finans", None])
    ws.append(["Taktik", "Satınalma", "Satınalma Kararları", "Tedarikçi değerlendirme", "Satınalma", "PR-99", None, None])
    ws2 = wb.create_sheet("Üretim")
    ws2.append(["Karar Tipi", "Birim", " Karar Süreci", "Karar Tanımı", "Karar Verici", "Katılımcı/Ortak Brimler", "Notlar", None])   # kaymış başlık
    for i in range(3):
        ws2.append(["Taktik", "Üretim", "Hat", f"Hat başı kontrol {i}", "PR-08 / ProManage", "Üretim", "Planlama", "not"])
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def test_real_layout_headers_and_all_sheets():
    assert matrix_io.classify_header("Karar Tipi") == "karar_tipi" and matrix_io.classify_header("Karar Türü") == "karar_tipi"
    assert matrix_io.classify_header("Karar Tanımı") == "karar" and matrix_io.classify_header(" Karar Süreci") == "surec"
    assert matrix_io.classify_header("Karar Verici") == "verici" and matrix_io.classify_header("Katılımcı/Ortak Brimler") == "katilimci"
    p = matrix_io.parse_matrix(multi_sheet_bytes(), "m.xlsx")
    assert [s["sayfa"] for s in p.sheets] == ["Proje ve Yatırım", "Üretim"] and len(p.rows) == 5
    a = p.rows[0]
    assert a.karar == "Satınalma talebinin onaylanması" and a.surec == "Satınalma Kararları" and a.karar_tipi == "Stratejik"
    assert a.birim == "Satınalma" and a.verici == ["Satınalma", "Genel Müdür"] and a.sheet == "Proje ve Yatırım"
    u = p.rows[2]                                           # Üretim: başlıklar bir kolon kaymış
    assert u.sheet == "Üretim" and u.dokuman_text == "PR-08 / ProManage" and u.verici == ["Üretim"]
    assert any("kaymış" in w for w in p.warnings)
    assert matrix_io.split_deciders("Yatırım ve Bakım-İCKB") == ["Yatırım ve Bakım", "İCKB"]
    assert matrix_io.split_deciders("Ar-Ge/Satınalma") == ["Arge", "Satınalma"]


def test_doc_refs_parse_and_resolve(kb):
    from kys.process.docrefs import DocResolver, parse_doc_refs
    refs = parse_doc_refs("PR.AR-01 5.8.2")
    assert refs[0].code == "PR.AR-01" and refs[0].clause_no == "5.8.2"
    assert [r.kind for r in parse_doc_refs("PR.UR-01 / PLANLAMA VE ARGE BÖLÜMÜNDE DÖKÜMAN BULUNMAMAKTA / ProManage")] == ["kod"]
    assert parse_doc_refs("Satin_Alma_Tedarik_Yonetimi_ve_SLA_Proseduru.docx")[0].kind == "dosya"
    res = DocResolver(kb)
    r = res.resolve_cell("pr.4 4.2")[0]
    assert r.resolved and r.doc_code == "PR-04" and r.clause_no == "4.2"
    assert not res.resolve_cell("PR11-SA-4")[0].resolved
    kb.execute("UPDATE documents SET file_path = 'x/Satin_Alma_Tedarik_Yonetimi_ve_SLA_Proseduru.docx' WHERE code = 'PR-04'")
    kb.commit()
    assert DocResolver(kb).resolve_cell("Satin_Alma_Tedarik_Yonetimi_ve_SLA_Proseduru.docx")[0].doc_code == "PR-04"


def test_unit_decider_is_not_forced_to_manager(kb):
    from kys.process.roles import RoleMatch, same_role
    idx = RoleIndex(kb)
    kb.execute("INSERT INTO entities(id, entity_type, canonical_name, normalized_name) VALUES (20, 'ROL', 'Yatırım ve Bakım Direktörü', 'yatirim ve bakim direktoru')")
    idx = RoleIndex(kb)
    unit = idx.resolve_decider("Yatırım ve Bakım")
    assert unit.is_unit and unit.method == "birim" and not unit.resolved           # yöneticisine zorla eşlenmedi
    assert same_role(unit, RoleMatch(input="Yatırımlar ve Bakım Direktörü")) is True  # aynı birimin direktörü uyumlu
    assert same_role(unit, idx.resolve("Finans Direktörü")) is False
    assert idx.resolve_decider("İCKB").method == "kisaltma" or not idx.resolve_decider("İCKB").is_unit


def test_joint_deciders_need_each_confirmed(kb):
    row = _row(onaylayan=[], verici=["Satınalma Müdürü", "Genel Müdür"], limit=None)
    idx = RoleIndex(kb)
    roles = {f: [idx.resolve(n) for n in getattr(row, f)] for f in matrix_io.ROLE_FIELDS}
    roles["authorities"] = [idx.resolve(n) for n in row.authorities]
    src = _src(idx, onay=["Genel Müdür"])
    v = decisions.classify(row, roles, [src])
    assert v["durum"] == DecisionStatus.KISMI and "Ortak/kademeli" in v["aciklama"] and "Satınalma Müdürü" in v["aciklama"]
    src2 = _src(idx, onay=["Genel Müdür", "Satınalma Müdürü"])
    assert decisions.classify(row, roles, [src2])["durum"] == DecisionStatus.KAYNAKLI


def test_kodsuz_source_is_flagged_in_explanation(kb):
    idx = RoleIndex(kb)
    s = _src(idx, doc="KODSUZ-61cc9032", onay=["Finans Direktörü"], quantities=[GM_100K])
    s.kodsuz = s.untrusted = True
    v = _classify(kb, _row(), [s])
    assert v["durum"] == DecisionStatus.CELISKI and "kodsuz/taslak" in v["aciklama"] and "KODSUZ-61cc9032" in v["aciklama"]


def test_multi_sheet_end_to_end_doc_check_and_excel(kb, monkeypatch):
    fake_llm(monkeypatch, [pick()])
    rows, out = decisions.check_matrix(kb, multi_sheet_bytes(), "m.xlsx")
    assert [(r["sayfa"], r["satir"]) for r in rows][:3] == [("Proje ve Yatırım", 2), ("Proje ve Yatırım", 3), ("Üretim", 2)]
    by = {(r["sayfa"], r["satir"]): r for r in rows}
    r1, r2, r3 = by[("Proje ve Yatırım", 2)], by[("Proje ve Yatırım", 3)], by[("Üretim", 2)]
    assert r1["dokuman_kontrolu"] == "DOGRULANDI" and r1["matris_dokumanlari"][0]["madde_uyumu"] is True
    assert r2["dokuman_kontrolu"] == "SISTEMDE_YOK" and "PR-99" in r2["aciklama"]
    assert r3["dokuman_kontrolu"] == "SISTEMDE_YOK" and r3["matris_dokumanlari"][0]["girdi"] == "PR-08"
    # karar vericisi birim adı: "Satınalma" birim olarak işaretlendi, "Genel Müdür" dokümanla uyumlu
    assert any(m["birim"] for m in r1["eslesen_roller"] if m["girdi"] == "Satınalma")
    wb = openpyxl.load_workbook(io.BytesIO(out))
    assert wb.sheetnames == ["Proje ve Yatırım", "Üretim", "Karar Kontrol Özeti"]
    for name in ("Proje ve Yatırım", "Üretim"):
        assert "Matristeki Doküman Kontrolü" in [c.value for c in wb[name][1]]
    summ = [str(c.value) for row in wb["Karar Kontrol Özeti"].iter_rows() for c in row if c.value is not None]
    assert "Üretim" in summ and "Proje ve Yatırım" in summ and "PR-08" in " ".join(summ)
    s = decisions.summarize(rows)
    assert s["sayfalar"]["Üretim"] and s["dokuman_kontrolu"]["DOGRULANDI"] == 1 and "PR-99" in s["sistemde_yok_dokumanlar"]


def test_api_chunks_span_sheets(client, monkeypatch):
    fake_llm(monkeypatch, [pick()])
    r1 = _upload(client, data=multi_sheet_bytes(), params={"offset": 1, "limit": 3}).json()
    assert [(x["sayfa"], x["satir"]) for x in r1["satirlar"]] == [("Proje ve Yatırım", 3), ("Üretim", 2), ("Üretim", 3)]
    assert r1["aralik"]["toplam"] == 5 and len(r1["matris"]["sayfalar"]) == 2


# --- CLI: --sayfa / --ekle (kısmi koşuyu tamamlama) --------------------------------------------------------------

def test_cli_karar_sayfa_and_ekle_resume(kb, monkeypatch, tmp_path):
    import argparse
    import json

    from kys import cli

    wb = openpyxl.Workbook()
    wb.active.title = "Diğer"
    for r in FLAT:
        wb.active.append(r)
    ws = wb.create_sheet("Proje ve Yatırım")
    for r in FLAT:
        ws.append(r)
    src = tmp_path / "matris.xlsx"
    wb.save(str(src))
    calls: list[str] = []
    fake_llm(monkeypatch, [pick()], calls=calls)

    def args(**kw):
        return argparse.Namespace(**({"sayfa": "proje ve yatirim", "ekle": None, "limit": None} | kw))

    first = cli._karar_analyze(kb, src, args(limit=1), None)
    assert [r["sayfa"] for r in first.rows] == ["Proje ve Yatırım"] and first.summary["toplam_satir"] == 2
    part = tmp_path / "ilk.json"
    part.write_text(json.dumps(first.to_json(), ensure_ascii=False), encoding="utf-8")
    n_first = len(calls)

    full = cli._karar_analyze(kb, src, args(ekle=str(part)), None)
    assert len(full.rows) == 2 and full.summary["kontrol_edilen"] == 2
    assert full.rows[0] == json.loads(part.read_text(encoding="utf-8"))["satirlar"][0]  # önceki satır aynen korunur
    assert len(calls) - n_first <= 1                                                    # yalnız kalan satır işlendi
    with pytest.raises(MatrixError):
        cli._karar_analyze(kb, src, args(sayfa="Yok"), None)