"""Rol birleştirme (Faz 3 ilk adımı): aynı rolün yazım varyantlarını gruplar, kullanıcı onayıyla tek varlıkta toplar.

Çıkarım her yeni yazımı ayrı bir `entities` satırı olarak açar ("Mekanik Bakım Teknisyeni", "Mekanik Teknisyeni" ...).
Akış:
1. `suggest_groups` — `roles.score_names` ile aday gruplar (LLM yok). Bir ad gruba ancak grubun TÜM üyelerine eşik
   üstü benziyorsa girer; zincirleme birleşme olmaz (Genel Müdür ≁ Genel Müdür Yardımcısı).
2. `render_xlsx` → kullanıcı KARAR (EVET/HAYIR) ve gerekirse ANA AD doldurur → `read_decisions`.
3. `save_decisions` kararları `data/kys_role_merges.json`'a yazar (DB yeniden kurulsa da korunur);
   `apply_saved` onaylı grupları DB'de birleştirir: eski adlar ana rolün takma adı olur, tüm atıflar ana role taşınır.
   `RoleIndex` takma adları okuduğu için karar kontrolü birleşmeyi kod değişikliği olmadan kullanır.
"""

import json
import sqlite3
from collections import defaultdict
from dataclasses import dataclass, field
from io import BytesIO

from .. import config
from ..ontology import EntityType, MergeDecision, NodeKind, ReviewStatus
from ..textnorm import fold
from .roles import SAME_ROLE_MIN, score_names

# `entities`'e işaret eden kolonlar (edges ve decisions.approval_chain ayrıca ele alınır)
_FK_COLUMNS = (("rules", "subject_entity_id"), ("rules", "object_entity_id"), ("processes", "owner_entity_id"),
               ("process_steps", "responsible_entity_id"), ("process_steps", "approver_entity_id"),
               ("decisions", "owner_entity_id"))

HEADERS = ("Grup", "Tip", "Önerilen ana ad", "Varyant", "Kullanım", "Skor", "KARAR", "ANA AD")


@dataclass
class Member:
    entity_id: int
    name: str
    usage: int
    score: float = 1.0   # ana ada benzerlik


@dataclass
class Group:
    entity_type: str
    members: list[Member]   # ilk üye önerilen ana ad (en çok kullanılan)

    @property
    def canonical(self) -> Member:
        return self.members[0]


@dataclass
class Decision:
    entity_type: str
    ana: str
    evet: list[str] = field(default_factory=list)
    hayir: list[str] = field(default_factory=list)


# --- Kalıcı kararlar (JSON) ---------------------------------------------------

def load() -> dict:
    p = config.settings.role_merges_path
    return json.loads(p.read_text(encoding="utf-8")) if p.exists() else {}


def _write(data: dict) -> None:
    p = config.settings.role_merges_path
    p.parent.mkdir(parents=True, exist_ok=True)
    p.write_text(json.dumps(data, ensure_ascii=False, indent=2, sort_keys=True), encoding="utf-8")


def _pair(a: str, b: str) -> list[str]:
    return sorted((fold(a), fold(b)))


def save_decisions(decisions: list[Decision]) -> dict:
    """Kararları JSON'a ekler. Ortak ad taşıyan önceki birleştirmeyle birleşir (ana ad yenisiyle değişir)."""
    data = load()
    for d in decisions:
        slot = data.setdefault(d.entity_type, {"birlesik": [], "ayri": []})
        names = [n for n in dict.fromkeys([d.ana, *d.evet]) if n]
        if len({fold(n) for n in names}) >= 2:
            folds = {fold(n) for n in names}
            keep, rest = [], []
            for g in slot["birlesik"]:
                (rest if folds & {fold(n) for n in [g["ana"], *g["varyantlar"]]} else keep).append(g)
            merged = list(dict.fromkeys(v for g in rest for v in [g["ana"], *g["varyantlar"]]))
            variants = [n for n in dict.fromkeys([*merged, *names]) if fold(n) != fold(d.ana)]
            keep.append({"ana": d.ana, "varyantlar": variants})
            slot["birlesik"] = keep
        for no in d.hayir:
            for yes in names:
                p = _pair(no, yes)
                if p[0] != p[1] and p not in slot["ayri"]:
                    slot["ayri"].append(p)
    _write(data)
    return data


# --- Gruplama ------------------------------------------------------------------

def usage_counts(conn: sqlite3.Connection) -> dict[int, int]:
    """Varlık başına atıf sayısı (kural, süreç, adım, karar, bilgi grafiği kenarı)."""
    counts: dict[int, int] = defaultdict(int)
    for table, col in _FK_COLUMNS:
        for r in conn.execute(f"SELECT {col} AS eid, COUNT(*) AS n FROM {table} WHERE {col} IS NOT NULL GROUP BY {col}"):
            counts[r["eid"]] += r["n"]
    for end in ("src", "dst"):
        for r in conn.execute(f"SELECT {end}_id AS eid, COUNT(*) AS n FROM edges WHERE {end}_kind = ? GROUP BY {end}_id",
                              (NodeKind.VARLIK.value,)):
            counts[r["eid"]] += r["n"]
    return counts


def suggest_groups(conn: sqlite3.Connection, entity_type: str = EntityType.ROL.value,
                   threshold: float = SAME_ROLE_MIN) -> list[Group]:
    """Aday gruplar (en az iki üyeli). Kullanıcının daha önce 'ayrı' dediği çiftler aynı gruba girmez."""
    rejected = {tuple(p) for p in load().get(entity_type, {}).get("ayri", [])}
    usage = usage_counts(conn)
    ents = [Member(r["id"], r["canonical_name"], usage.get(r["id"], 0)) for r in conn.execute(
        "SELECT id, canonical_name FROM entities WHERE entity_type = ?", (entity_type,))]
    ents.sort(key=lambda m: (-m.usage, fold(m.name)))
    clusters: list[list[Member]] = []
    for e in ents:
        best, best_score = None, 0.0
        for cl in clusters:
            if any(tuple(_pair(e.name, m.name)) in rejected for m in cl):
                continue
            scores = [score_names(e.name, m.name) for m in cl]
            if min(scores) >= threshold and sum(scores) / len(scores) > best_score:
                best, best_score = cl, sum(scores) / len(scores)
        if best is not None:
            best.append(e)
        else:
            clusters.append([e])
    groups = []
    for cl in clusters:
        if len(cl) < 2:
            continue
        for m in cl[1:]:
            m.score = score_names(m.name, cl[0].name)
        groups.append(Group(entity_type, cl))
    groups.sort(key=lambda g: (-sum(m.usage for m in g.members), fold(g.canonical.name)))
    return groups


# --- Excel ---------------------------------------------------------------------

def render_xlsx(groups: list[Group]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    ws = wb.active
    ws.title = "Rol Grupları"
    ws.append(HEADERS)
    for c in ws[1]:
        c.font = Font(bold=True)
    shade = PatternFill("solid", fgColor="EEF3F8")
    for gi, g in enumerate(groups, 1):
        for m in g.members:
            ws.append([gi, g.entity_type, g.canonical.name, m.name, m.usage, round(m.score, 3), None, None])
            if gi % 2 == 0:
                for c in ws[ws.max_row]:
                    c.fill = shade
    dv = DataValidation(type="list", formula1=f'"{MergeDecision.EVET.value},{MergeDecision.HAYIR.value}"',
                        allow_blank=True)
    ws.add_data_validation(dv)
    if ws.max_row > 1:
        dv.add(f"G2:G{ws.max_row}")
    for col, width in zip("ABCDEFGH", (7, 8, 36, 36, 9, 7, 9, 36)):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    info = wb.create_sheet("Açıklama")
    for line in (
        "Her satır bir rol yazımıdır; aynı Grup numarasındaki satırlar aynı rol olarak önerilir.",
        f"KARAR: {MergeDecision.EVET.value} = bu yazım gruptaki rolle aynıdır (birleştirilir). "
        f"{MergeDecision.HAYIR.value} = ayrı roldür (bir daha bu grupta önerilmez). Boş = karar verilmedi, dokunulmaz.",
        "Önerilen ana ad satırına da EVET yazın; HAYIR yazarsanız o ad gruptan çıkar.",
        "ANA AD: grubun kalıcı adı. Boşsa 'Önerilen ana ad' kullanılır. Grubun herhangi bir satırına yazmanız yeterli.",
        "Sonra: kys rol-onayla <bu dosya> (önizleme) → kys rol-onayla <bu dosya> --apply",
    ):
        info.append([line])
    info.column_dimensions["A"].width = 120
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def read_decisions(data: bytes) -> list[Decision]:
    from openpyxl import load_workbook

    ws = load_workbook(BytesIO(data), read_only=True, data_only=True)["Rol Grupları"]
    rows = ws.iter_rows(values_only=True)
    header = [str(h or "").strip() for h in next(rows)]
    missing = [h for h in HEADERS if h not in header]
    if missing:
        raise ValueError(f"Eksik sütun: {', '.join(missing)}")
    idx = {h: header.index(h) for h in HEADERS}
    groups: dict[object, dict] = {}
    for row in rows:
        def val(h):
            v = row[idx[h]] if idx[h] < len(row) else None
            return " ".join(str(v).split()) if v is not None else ""
        if not val("Grup") or not val("Varyant"):
            continue
        g = groups.setdefault(val("Grup"), {"tip": val("Tip") or EntityType.ROL.value,
                                            "oneri": val("Önerilen ana ad"), "ana": "", "evet": [], "hayir": []})
        g["ana"] = g["ana"] or val("ANA AD")
        karar = fold(val("KARAR")).upper()
        if karar == MergeDecision.EVET.value:
            g["evet"].append(val("Varyant"))
        elif karar == MergeDecision.HAYIR.value:
            g["hayir"].append(val("Varyant"))
        elif karar:
            raise ValueError(f"Grup {val('Grup')}, '{val('Varyant')}': KARAR yalnızca "
                             f"{MergeDecision.EVET.value}/{MergeDecision.HAYIR.value} olabilir (yazılan: {val('KARAR')})")
    out = []
    for g in groups.values():
        if g["tip"] not in {t.value for t in EntityType}:
            raise ValueError(f"Bilinmeyen tip: {g['tip']}")
        if not g["evet"] and not g["hayir"]:
            continue
        ana = g["ana"] or (g["oneri"] if g["oneri"] not in g["hayir"] else (g["evet"][0] if g["evet"] else ""))
        out.append(Decision(g["tip"], ana, g["evet"], g["hayir"]))
    return out


# --- DB'de birleştirme -----------------------------------------------------------

def _entity_ids(conn, entity_type: str, names: list[str]) -> list[int]:
    folds = list({fold(n) for n in names if n})
    if not folds:
        return []
    marks = ",".join("?" * len(folds))
    rows = conn.execute(
        f"""SELECT DISTINCT e.id FROM entities e LEFT JOIN entity_aliases a ON a.entity_id = e.id
            WHERE e.entity_type = ? AND (e.normalized_name IN ({marks}) OR a.normalized_alias IN ({marks}))""",
        (entity_type, *folds, *folds))
    return [r["id"] for r in rows]


def merge_group(conn: sqlite3.Connection, entity_type: str, ana: str, names: list[str]) -> dict:
    """`ana` ve `names` adlı (ya da takma adlı) varlıkları tek varlıkta toplar. Commit etmez.
    Dönen sözlük: birlesen (silinen varlık), tasinan (taşınan atıf), takma_ad (eklenen takma ad)."""
    stats = {"birlesen": 0, "tasinan": 0, "takma_ad": 0}
    ids = _entity_ids(conn, entity_type, [ana, *names])
    if not ids:
        return stats
    ana_fold = fold(ana)
    by_name = {r["id"]: r for r in conn.execute(
        f"SELECT id, canonical_name, normalized_name FROM entities WHERE id IN ({','.join('?' * len(ids))})", ids)}
    keeper = next((i for i in ids if by_name[i]["normalized_name"] == ana_fold), None)
    if keeper is None:
        usage = usage_counts(conn)
        keeper = max(ids, key=lambda i: (usage.get(i, 0), -i))
    old_names = [by_name[i]["canonical_name"] for i in ids]
    for other in (i for i in ids if i != keeper):
        for table, col in _FK_COLUMNS:
            stats["tasinan"] += conn.execute(f"UPDATE {table} SET {col} = ? WHERE {col} = ?", (keeper, other)).rowcount
        for end in ("src", "dst"):
            stats["tasinan"] += conn.execute(
                f"UPDATE OR IGNORE edges SET {end}_id = ? WHERE {end}_kind = ? AND {end}_id = ?",
                (keeper, NodeKind.VARLIK.value, other)).rowcount
            conn.execute(f"DELETE FROM edges WHERE {end}_kind = ? AND {end}_id = ?", (NodeKind.VARLIK.value, other))
        for d in conn.execute("SELECT id, approval_chain FROM decisions WHERE approval_chain IS NOT NULL").fetchall():
            chain = json.loads(d["approval_chain"])
            if other in chain:
                conn.execute("UPDATE decisions SET approval_chain = ? WHERE id = ?",
                             (json.dumps([keeper if x == other else x for x in chain]), d["id"]))
                stats["tasinan"] += 1
        conn.execute("UPDATE OR IGNORE entity_aliases SET entity_id = ? WHERE entity_id = ?", (keeper, other))
        conn.execute("DELETE FROM entities WHERE id = ?", (other,))
        stats["birlesen"] += 1
    conn.execute("UPDATE entities SET canonical_name = ?, normalized_name = ?, review_status = ? WHERE id = ?",
                 (" ".join(ana.split()), ana_fold, ReviewStatus.ONAYLANDI.value, keeper))
    for n in dict.fromkeys([*old_names, *names]):
        if n and fold(n) != ana_fold:
            stats["takma_ad"] += conn.execute(
                "INSERT OR IGNORE INTO entity_aliases(entity_id, alias, normalized_alias) VALUES (?,?,?)",
                (keeper, n, fold(n))).rowcount
    conn.execute("DELETE FROM entity_aliases WHERE entity_id = ? AND normalized_alias = ?", (keeper, ana_fold))
    return stats


def apply_saved(conn: sqlite3.Connection, data: dict | None = None, commit: bool = True) -> dict:
    """JSON'daki tüm onaylı birleştirmeleri uygular (idempotent). `commit=False`: önizleme, geri alınır."""
    data = load() if data is None else data
    total = {"grup": 0, "birlesen": 0, "tasinan": 0, "takma_ad": 0}
    try:
        for entity_type, slot in data.items():
            for g in slot.get("birlesik", []):
                s = merge_group(conn, entity_type, g["ana"], g["varyantlar"])
                if s["birlesen"] or s["takma_ad"]:
                    total["grup"] += 1
                for k in ("birlesen", "tasinan", "takma_ad"):
                    total[k] += s[k]
    except Exception:
        conn.rollback()
        raise
    if commit:
        conn.commit()
    else:
        conn.rollback()
    return total
